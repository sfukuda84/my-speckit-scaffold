"""speckit-presentation の build_pptx.py のテスト。

解析と検査は標準ライブラリだけで確かめる。pptx の生成は python-pptx と PyYAML があるときだけ確かめる:
    uv run --no-project --with python-pptx --with pyyaml python -m unittest discover -s tests
"""

from __future__ import annotations

import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parents[2] / "skills" / "speckit" / "speckit-presentation" / "scripts"
TEMPLATES = SCRIPTS.parent / "templates"
sys.path.insert(0, str(SCRIPTS))

import build_pptx  # noqa: E402

HAS_PPTX = importlib.util.find_spec("pptx") is not None and importlib.util.find_spec("yaml") is not None

DECK = """<!-- layout: title -->
# 題名
副題

---

<!-- layout: bullets -->
# 箇条書きの見出し
- 一つ目 **太字**
  - 下位
- 二つ目
> 出典: docs/project.md §2

<!-- notes
ノートの 1 行目
ノートの 2 行目
-->

---

<!-- layout: two-column -->
# 比較
## 左
- a
## 右
- b

---

<!-- layout: table -->
# 表
| 機能 | 区分 |
|---|---|
| 予約 | MVP |

---

<!-- layout: stats -->
# 数字
- **18万件** 出願
- **90人** 利用者

---

<!-- layout: chart -->
<!-- chart: line -->
# グラフ
| 期 | 売上 | 費用 |
|---|---|---|
| 1年目 | 0 | 80,400 |
| 2年目 | 352,800 | 190,000 |

---

<!-- layout: section -->
# 章

---

<!-- layout: message -->
# お願いしたいこと
補足
"""


class ParseTest(unittest.TestCase):
    def test_parse_layouts_and_parts(self) -> None:
        deck = build_pptx.parse_deck(DECK)
        self.assertEqual([s.layout for s in deck.slides],
                         ["title", "bullets", "two-column", "table", "stats", "chart", "section", "message"])
        title, bullets, cols, table, stats, chart = deck.slides[:6]
        self.assertEqual(title.paragraphs, ["副題"])
        self.assertEqual([(b.text, b.level) for b in bullets.bullets],
                         [("一つ目 **太字**", 0), ("下位", 1), ("二つ目", 0)])
        self.assertEqual(bullets.sources, ["docs/project.md §2"])
        self.assertEqual(bullets.notes, "ノートの 1 行目\nノートの 2 行目")
        self.assertEqual([name for name, _ in cols.columns], ["左", "右"])
        self.assertEqual(table.table, [["機能", "区分"], ["予約", "MVP"]])
        self.assertEqual(len(stats.bullets), 2)
        self.assertEqual(chart.chart, "line")

    def test_display_width(self) -> None:
        self.assertEqual(build_pptx.display_width("日本語"), 3)
        self.assertEqual(build_pptx.display_width("abcd"), 2)
        self.assertEqual(build_pptx.display_width("**太字**a"), 2.5)

    def test_parse_number(self) -> None:
        self.assertEqual(build_pptx.parse_number("352,800"), 352800)
        self.assertEqual(build_pptx.parse_number("12%"), 12)
        self.assertIsNone(build_pptx.parse_number("未定"))


class CheckTest(unittest.TestCase):
    def check(self, text: str) -> tuple[list[str], list[str]]:
        deck = build_pptx.parse_deck(text)
        design = build_pptx.merge_design(build_pptx.DEFAULT_DESIGN, {})
        return build_pptx.check_deck(deck, design, Path("."))

    def test_valid_deck(self) -> None:
        errors, warnings = self.check(DECK)
        self.assertEqual(errors, [])
        self.assertEqual(warnings, [])

    def test_errors(self) -> None:
        errors, _ = self.check("<!-- layout: unknown -->\n# x\n")
        self.assertIn("未知のレイアウト", errors[0])
        errors, _ = self.check("<!-- layout: chart -->\n# x\n| a | b |\n|---|---|\n| 1年目 | 未定 |\n")
        self.assertTrue(any("数値でない" in e for e in errors))
        errors, _ = self.check("<!-- layout: stats -->\n# x\n- 太字なし\n")
        self.assertTrue(any("**数値** 説明" in e for e in errors))
        errors, _ = self.check("<!-- layout: image -->\n# x\n![a](missing.png)\n")
        self.assertTrue(any("見つからない" in e for e in errors))
        errors, _ = self.check("<!-- layout: two-column -->\n# x\n## 左\n- a\n")
        self.assertTrue(any("2 つの列" in e for e in errors))

    def test_limit_warnings(self) -> None:
        many = "\n".join(f"- 項目{i}" for i in range(8))
        _, warnings = self.check(f"# 見出し\n{many}\n")
        self.assertTrue(any("上限 6" in w for w in warnings))
        _, warnings = self.check("# 見出し\n- " + "長" * 60 + "\n")
        self.assertTrue(any("全角換算 60" in w for w in warnings))

    def test_template_example_is_valid(self) -> None:
        """同梱の slides.md の見本は、画像のスライドを除いてエラーにならない。"""
        text = (TEMPLATES / "slides.md").read_text(encoding="utf-8")
        _, body = text.split("\n---\n", 1)[0], text  # front matter は PyYAML が要るので外して検査する
        body = text[text.find("\n---\n", 4) + 5:]
        errors, _ = self.check(body)
        self.assertTrue(all("画像" in e for e in errors), errors)


@unittest.skipUnless(HAS_PPTX, "python-pptx と PyYAML がない")
class BuildTest(unittest.TestCase):
    def test_build_pptx(self) -> None:
        from pptx import Presentation
        from pptx.oxml.ns import qn

        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp) / "internal"
            d.mkdir()
            (d / "slides.md").write_text("---\ntitle: t\ndate: 2026-09-23\n---\n\n" + DECK, encoding="utf-8")
            (Path(tmp) / "design.yaml").write_text("fonts:\n  east_asian: Yu Gothic\n", encoding="utf-8")
            self.assertEqual(build_pptx.main([str(d / "slides.md")]), 0)
            prs = Presentation(str(d / "proposal.pptx"))
            self.assertEqual(len(prs.slides), 8)
            self.assertAlmostEqual(prs.slide_width / 914400, 13.333, places=2)
            # 和文フォントの指定（a:ea）が入っている
            ea = prs.slides[1].shapes._spTree.xpath(".//a:ea")
            self.assertTrue(ea and all(e.get("typeface") == "Yu Gothic" for e in ea))
            # 発表者ノート、出典、グラフ、表
            self.assertEqual(prs.slides[1].notes_slide.notes_text_frame.text, "ノートの 1 行目\nノートの 2 行目")
            texts = [sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame]
            self.assertTrue(any("出典: docs/project.md §2" in t for t in texts))
            charts = [sh for sh in prs.slides[5].shapes if sh.has_chart]
            self.assertEqual(len(charts), 1)
            self.assertEqual([s.name for s in charts[0].chart.series], ["売上", "費用"])
            tables = [sh for sh in prs.slides[3].shapes if sh.has_table]
            self.assertEqual(tables[0].table.cell(1, 0).text, "予約")
            # 太字の run が分かれている
            runs = [r for sh in prs.slides[1].shapes if sh.has_text_frame
                    for p in sh.text_frame.paragraphs for r in p.runs]
            self.assertTrue(any(r.text == "太字" and r.font.bold for r in runs))
            self.assertIsNotNone(runs[0]._r.find(qn("a:rPr")))

    def test_stat_font_shrinks_and_table_columns_follow_content(self) -> None:
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp) / "internal"
            d.mkdir()
            (d / "slides.md").write_text(
                "<!-- layout: stats -->\n# 数字\n- **約 27.3 万件/年** 市場\n- **90人** 利用者\n- **3.6万件** 出願\n\n---\n\n"
                "<!-- layout: table -->\n# 表\n| # | マイルストーン | 時期 |\n|---|---|---|\n| M0 | 商標データの申込の見通し | 2026 年 10〜12 月 |\n",
                encoding="utf-8")
            self.assertEqual(build_pptx.main([str(d / "slides.md")]), 0)
            prs = Presentation(str(d / "proposal.pptx"))
            sizes = {r.text: r.font.size.pt for sh in prs.slides[0].shapes if sh.has_text_frame
                     for p in sh.text_frame.paragraphs for r in p.runs}
            self.assertLess(sizes["約 27.3 万件/年"], sizes["90人"])  # 長い数値だけ小さくなる
            self.assertEqual(sizes["90人"], build_pptx.DEFAULT_DESIGN["sizes"]["stat"])
            table = [sh for sh in prs.slides[1].shapes if sh.has_table][0].table
            self.assertLess(table.columns[0].width, table.columns[1].width)

    def test_template_pptx_is_used(self) -> None:
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as tmp:
            base = Presentation()
            base.slides.add_slide(base.slide_layouts[0])  # 雛形に元からあるスライドは取り除かれる
            base.save(str(Path(tmp) / "brand.pptx"))
            (Path(tmp) / "design.yaml").write_text("template: brand.pptx\ntemplate_layout: Blank\n", encoding="utf-8")
            d = Path(tmp) / "customer"
            d.mkdir()
            (d / "slides.md").write_text("# 見出し\n- 項目\n", encoding="utf-8")
            self.assertEqual(build_pptx.main([str(d / "slides.md")]), 0)
            prs = Presentation(str(d / "proposal.pptx"))
            self.assertEqual(len(prs.slides), 1)
            self.assertEqual(prs.slides[0].slide_layout.name, "Blank")


if __name__ == "__main__":
    unittest.main()
