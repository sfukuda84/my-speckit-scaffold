# /// script
# requires-python = ">=3.9"
# dependencies = ["python-pptx>=1.0", "PyYAML>=6"]
# ///
"""slides.md（スライドの内容）と design.yaml（デザインスペック）から .pptx を作る。

使い方:
    uv run build_pptx.py <slides.md> [--design <design.yaml>] [-o <出力.pptx>]
    uv run build_pptx.py <slides.md> --check        # 変換せずに、書式とデザインの上限だけを確かめる

    --design を省くと、slides.md の front matter の design、なければ slides.md の 1 つ上の design.yaml を使う。
    -o を省くと、slides.md と同じディレクトリの proposal.pptx に書き出す。
    uv がない環境では `pip install python-pptx PyYAML` のうえで `python3 build_pptx.py …` を実行する。

終了コード: エラーがあれば 1、なければ 0（上限を超えた警告は終了コードに影響しない。--strict で 1 にする）。
"""

from __future__ import annotations

import re
import sys
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

LAYOUTS = ("title", "section", "bullets", "two-column", "table", "stats", "chart", "image", "message")
CHART_TYPES = ("bar", "column", "line", "pie")
DIRECTIVE_RE = re.compile(r"<!--\s*(layout|chart)\s*:\s*([\w-]+)\s*-->")
NOTES_RE = re.compile(r"<!--\s*notes\s*(.*?)-->", re.DOTALL)
IMAGE_RE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")
BULLET_RE = re.compile(r"^( *)[-*] (.*)$")
STAT_RE = re.compile(r"^\*\*(.+?)\*\*\s*(.*)$")

DEFAULT_DESIGN = {
    "slide": {"width_in": 13.333, "height_in": 7.5},
    "template": None,
    "template_layout": None,
    "fonts": {"latin": "Arial", "east_asian": "Meiryo"},
    "colors": {"primary": "1F3A5F", "accent": "E07A1F", "text": "222222", "muted": "6B7280",
               "background": "FFFFFF", "light": "F3F4F6", "on_primary": "FFFFFF"},
    "sizes": {"title": 40, "heading": 28, "body": 18, "small": 11, "stat": 44, "message": 36},
    "margins_in": {"left": 0.6, "right": 0.6, "top": 0.45, "bottom": 0.45},
    "footer": {"text": "", "page_number": True},
    "logo": None,
    "limits": {
        "title_max_chars": 36,
        "bullets": {"max_items": 6, "max_chars": 45},
        "two-column": {"max_items": 5, "max_chars": 22},
        "table": {"max_rows": 8, "max_cols": 5, "max_chars": 18},
        "stats": {"max_items": 4, "max_chars": 24, "max_value_chars": 8},
        "message": {"max_chars": 40},
    },
}


# --- スライドの解析（python-pptx に依存しない） ----------------------------------------

@dataclass
class Bullet:
    text: str
    level: int = 0


@dataclass
class Slide:
    number: int
    layout: str = "bullets"
    chart: str | None = None
    title: str = ""
    paragraphs: list[str] = field(default_factory=list)
    bullets: list[Bullet] = field(default_factory=list)
    columns: list[tuple[str, list[Bullet]]] = field(default_factory=list)
    table: list[list[str]] = field(default_factory=list)
    image: tuple[str, str] | None = None
    sources: list[str] = field(default_factory=list)
    notes: str = ""


@dataclass
class Deck:
    meta: dict
    slides: list[Slide]


class BuildError(Exception):
    pass


def display_width(text: str) -> float:
    """全角を 1、半角を 0.5 として数えた文字数。"""
    plain = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    return sum(1.0 if unicodedata.east_asian_width(ch) in ("F", "W", "A") else 0.5 for ch in plain)


def split_front_matter(text: str) -> tuple[dict, str]:
    text = text.replace("\r\n", "\n")
    if not text.startswith("---\n"):
        return {}, text
    end = text.find("\n---\n", 4)
    if end < 0:
        raise BuildError("front matter の終わりの --- がない")
    import yaml  # 遅延 import（解析だけのテストで PyYAML を要らなくするため、front matter がある場合だけ使う）

    meta = yaml.safe_load(text[4:end]) or {}
    if not isinstance(meta, dict):
        raise BuildError("front matter は key: value の形にする")
    return meta, text[end + 5:]


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
            continue
        rows.append(cells)
    return rows


def parse_slide(number: int, chunk: str) -> Slide:
    slide = Slide(number=number)
    notes = NOTES_RE.findall(chunk)
    slide.notes = "\n".join(n.strip() for n in notes)
    chunk = NOTES_RE.sub("", chunk)
    for key, value in DIRECTIVE_RE.findall(chunk):
        if key == "layout":
            slide.layout = value
        else:
            slide.chart = value
    chunk = DIRECTIVE_RE.sub("", chunk)

    table_lines: list[str] = []
    current_column: tuple[str, list[Bullet]] | None = None
    for raw in chunk.split("\n"):
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("# "):
            slide.title = line[2:].strip()
        elif line.startswith("## "):
            current_column = (line[3:].strip(), [])
            slide.columns.append(current_column)
        elif line.lstrip().startswith(">"):
            body = line.lstrip()[1:].strip()
            slide.sources.append(re.sub(r"^出典\s*[:：]\s*", "", body))
        elif line.lstrip().startswith("|"):
            table_lines.append(line)
        elif IMAGE_RE.match(line.strip()):
            m = IMAGE_RE.match(line.strip())
            slide.image = (m.group(1), m.group(2))
        elif BULLET_RE.match(line):
            m = BULLET_RE.match(line)
            bullet = Bullet(text=m.group(2).strip(), level=min(len(m.group(1)) // 2, 1))
            (current_column[1] if current_column else slide.bullets).append(bullet)
        else:
            slide.paragraphs.append(line.strip())
    slide.table = parse_table(table_lines)
    return slide


def parse_deck(text: str) -> Deck:
    meta, body = split_front_matter(text)
    chunks = re.split(r"^---\s*$", body, flags=re.MULTILINE)
    slides = [parse_slide(i, c) for i, c in enumerate((c for c in chunks if c.strip()), start=1)]
    if not slides:
        raise BuildError("スライドが 1 枚もない")
    return Deck(meta=meta, slides=slides)


# --- 検査 -------------------------------------------------------------------------

def merge_design(base: dict, override: dict) -> dict:
    result = dict(base)
    for key, value in (override or {}).items():
        if isinstance(value, dict) and isinstance(result.get(key), dict):
            result[key] = merge_design(result[key], value)
        else:
            result[key] = value
    return result


def check_deck(deck: Deck, design: dict, base_dir: Path) -> tuple[list[str], list[str]]:
    """(エラー, 警告) を返す。"""
    errors: list[str] = []
    warnings: list[str] = []
    limits = design["limits"]

    def over(where: str, text: str, max_chars: float) -> None:
        width = display_width(text)
        if width > max_chars:
            warnings.append(f"{where}: 全角換算 {width:g} 文字で、上限 {max_chars} を超えている（「{text[:20]}…」）")

    for s in deck.slides:
        where = f"スライド {s.number}（{s.layout}）"
        if s.layout not in LAYOUTS:
            errors.append(f"{where}: 未知のレイアウト。{', '.join(LAYOUTS)} のいずれかにする")
            continue
        if not s.title and s.layout != "image":
            errors.append(f"{where}: 見出し（# …）がない")
        over(f"{where} の見出し", s.title, limits["title_max_chars"])
        if s.layout == "bullets":
            lim = limits["bullets"]
            if not s.bullets and not s.paragraphs:
                errors.append(f"{where}: 箇条書きがない")
            if len(s.bullets) > lim["max_items"]:
                warnings.append(f"{where}: 箇条書きが {len(s.bullets)} 項目で、上限 {lim['max_items']} を超えている")
            for b in s.bullets:
                over(where, b.text, lim["max_chars"])
        elif s.layout == "two-column":
            lim = limits["two-column"]
            if len(s.columns) != 2:
                errors.append(f"{where}: ## の小見出しで 2 つの列を書く（今は {len(s.columns)} 列）")
            for name, items in s.columns:
                if len(items) > lim["max_items"]:
                    warnings.append(f"{where}: 「{name}」が {len(items)} 項目で、上限 {lim['max_items']} を超えている")
                for b in items:
                    over(where, b.text, lim["max_chars"])
        elif s.layout in ("table", "chart"):
            if len(s.table) < 2:
                errors.append(f"{where}: 見出し行と 1 行以上のデータ行を持つ表を書く")
            if s.layout == "table":
                lim = limits["table"]
                if len(s.table) - 1 > lim["max_rows"]:
                    warnings.append(f"{where}: 表が {len(s.table) - 1} 行で、上限 {lim['max_rows']} を超えている")
                if s.table and len(s.table[0]) > lim["max_cols"]:
                    warnings.append(f"{where}: 表が {len(s.table[0])} 列で、上限 {lim['max_cols']} を超えている")
                for row in s.table:
                    for cell in row:
                        over(where, cell, lim["max_chars"])
            else:
                if (s.chart or "column") not in CHART_TYPES:
                    errors.append(f"{where}: chart は {', '.join(CHART_TYPES)} のいずれかにする")
                for row in s.table[1:]:
                    for cell in row[1:]:
                        if parse_number(cell) is None:
                            errors.append(f"{where}: グラフの値「{cell}」が数値でない")
        elif s.layout == "stats":
            lim = limits["stats"]
            if not s.bullets:
                errors.append(f"{where}: `- **数値** 説明` の形で項目を書く")
            if len(s.bullets) > lim["max_items"]:
                warnings.append(f"{where}: 数値が {len(s.bullets)} 個で、上限 {lim['max_items']} を超えている")
            for b in s.bullets:
                if not STAT_RE.match(b.text):
                    errors.append(f"{where}: 「{b.text}」は `**数値** 説明` の形にする")
                else:
                    over(where, STAT_RE.match(b.text).group(2), lim["max_chars"])
                    over(f"{where} の数値", STAT_RE.match(b.text).group(1), lim.get("max_value_chars", 8))
        elif s.layout == "image":
            if not s.image:
                errors.append(f"{where}: ![説明](パス) の画像がない")
            elif not (base_dir / s.image[1]).is_file():
                errors.append(f"{where}: 画像 {s.image[1]} が見つからない")
        elif s.layout == "message":
            over(where, s.title, limits["message"]["max_chars"])
    return errors, warnings


def parse_number(text: str) -> float | None:
    cleaned = re.sub(r"[,\s円%％人件社店舗]", "", text)
    try:
        return float(cleaned)
    except ValueError:
        return None


# --- pptx の生成 ------------------------------------------------------------------

class Renderer:
    def __init__(self, design: dict, base_dir: Path, design_dir: Path, meta: dict) -> None:
        from pptx import Presentation
        from pptx.util import Inches

        self.d = design
        self.meta = meta
        self.base_dir = base_dir
        self.design_dir = design_dir
        self.Inches = Inches
        template = design.get("template")
        self.use_template = bool(template)
        if template:
            path = (design_dir / template).resolve()
            if not path.is_file():
                raise BuildError(f"雛形の pptx {path} がない")
            self.prs = Presentation(str(path))
            self._drop_existing_slides()
        else:
            self.prs = Presentation()
            self.prs.slide_width = Inches(design["slide"]["width_in"])
            self.prs.slide_height = Inches(design["slide"]["height_in"])
        self.layout = self._pick_layout()
        self.W = self.prs.slide_width
        self.H = self.prs.slide_height
        m = design["margins_in"]
        self.left, self.right = Inches(m["left"]), Inches(m["right"])
        self.top, self.bottom = Inches(m["top"]), Inches(m["bottom"])

    # --- 雛形 ---
    def _drop_existing_slides(self) -> None:
        sld_ids = self.prs.slides._sldIdLst
        for sld_id in list(sld_ids):
            self.prs.part.drop_rel(sld_id.rId)
            sld_ids.remove(sld_id)

    def _pick_layout(self):
        layouts = list(self.prs.slide_layouts)
        name = self.d.get("template_layout")
        if name:
            for layout in layouts:
                if layout.name == name:
                    return layout
            raise BuildError(f"雛形にレイアウト「{name}」がない（ある: {', '.join(l.name for l in layouts)}）")
        return min(layouts, key=lambda l: len(l.placeholders))

    # --- 文字 ---
    def color(self, key: str):
        from pptx.dml.color import RGBColor

        return RGBColor.from_string(str(self.d["colors"].get(key, key)).lstrip("#"))

    def style_run(self, run, size: float, color: str = "text", bold: bool = False) -> None:
        from lxml import etree
        from pptx.oxml.ns import qn
        from pptx.util import Pt

        font = run.font
        font.size = Pt(size)
        font.bold = bold
        font.color.rgb = self.color(color)
        font.name = self.d["fonts"]["latin"]
        # 和文フォント（東アジア用）は python-pptx では指定できないため、a:latin の後ろに a:ea と a:cs を足す
        rpr = run._r.get_or_add_rPr()
        latin = rpr.find(qn("a:latin"))
        anchor = latin
        for tag in ("a:ea", "a:cs"):
            el = rpr.find(qn(tag))
            if el is None:
                el = etree.SubElement(rpr, qn(tag))
                anchor.addnext(el)
            el.set("typeface", self.d["fonts"]["east_asian"])
            anchor = el

    def add_runs(self, paragraph, text: str, size: float, color: str = "text", bold: bool = False) -> None:
        """**太字** を含む文字列を、太字と通常の run に分けて足す。"""
        for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
            if not part:
                continue
            run = paragraph.add_run()
            run.text = part
            self.style_run(run, size, color, bold or i % 2 == 1)

    def text_box(self, x, y, w, h, *, autofit: bool = True, anchor: str = "top"):
        from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE

        box = self.slide.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        if autofit:
            tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        return tf

    def write(self, tf, text: str, size: float, color: str = "text", bold: bool = False, align=None, first=True):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        if align is not None:
            p.alignment = align
        self.add_runs(p, text, size, color, bold)
        return p

    def rect(self, x, y, w, h, color: str):
        from pptx.enum.shapes import MSO_SHAPE

        shape = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = self.color(color)
        shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    # --- 共通の枠 ---
    def new_slide(self, dark: bool = False) -> None:
        self.slide = self.prs.slides.add_slide(self.layout)
        for ph in list(self.slide.placeholders):
            ph._element.getparent().remove(ph._element)
        if not self.use_template:
            self.rect(0, 0, self.W, self.H, "primary" if dark else "background")

    def content_width(self):
        return self.W - self.left - self.right

    def heading(self, s: Slide):
        from pptx.util import Inches

        tf = self.text_box(self.left, self.top, self.content_width(), Inches(0.9), anchor="middle")
        self.write(tf, s.title, self.d["sizes"]["heading"], "primary", bold=True)
        self.rect(self.left, self.top + Inches(0.95), Inches(1.2), Inches(0.06), "accent")
        return self.top + Inches(1.2)

    def footer(self, s: Slide, page: int, dark: bool = False) -> None:
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches

        y = self.H - self.bottom - Inches(0.3)
        muted = "on_primary" if dark else "muted"
        if s.sources:
            tf = self.text_box(self.left, y - Inches(0.35), self.content_width() - Inches(1.2), Inches(0.6),
                               anchor="bottom")
            self.write(tf, "出典: " + " / ".join(s.sources), self.d["sizes"]["small"], muted)
        footer = self.d.get("footer") or {}
        if footer.get("page_number", True) and page > 1:
            tf = self.text_box(self.W - self.right - Inches(1.0), y, Inches(1.0), Inches(0.3), autofit=False)
            self.write(tf, str(page), self.d["sizes"]["small"], muted, align=PP_ALIGN.RIGHT)
        if footer.get("text") and page > 1:
            tf = self.text_box(self.left, y + Inches(0.05), Inches(6), Inches(0.3), autofit=False)
            self.write(tf, footer["text"], self.d["sizes"]["small"], muted)
        logo = self.d.get("logo")
        if logo and page > 1:
            path = self.design_dir / logo
            if path.is_file():
                self.slide.shapes.add_picture(str(path), self.W - self.right - Inches(1.6), self.top,
                                              height=Inches(0.5))

    def notes(self, s: Slide) -> None:
        if s.notes:
            self.slide.notes_slide.notes_text_frame.text = s.notes

    # --- レイアウト ---
    def render(self, deck: Deck) -> None:
        for page, s in enumerate(deck.slides, start=1):
            getattr(self, "layout_" + s.layout.replace("-", "_"))(s, page)
            self.notes(s)

    def layout_title(self, s: Slide, page: int) -> None:
        from pptx.util import Inches

        self.new_slide(dark=True)
        tf = self.text_box(self.left, Inches(2.0), self.content_width(), Inches(1.6), anchor="bottom")
        self.write(tf, s.title, self.d["sizes"]["title"], "on_primary", bold=True)
        self.rect(self.left, Inches(3.75), Inches(1.6), Inches(0.08), "accent")
        tf = self.text_box(self.left, Inches(4.0), self.content_width(), Inches(1.4))
        for i, text in enumerate(s.paragraphs):
            self.write(tf, text, self.d["sizes"]["body"] + 4, "on_primary", first=(i == 0))
        meta = " ｜ ".join(str(v) for v in (self.meta.get("date"), self.meta.get("author")) if v)
        if meta:
            tf = self.text_box(self.left, self.H - self.bottom - Inches(0.8), self.content_width(), Inches(0.5))
            self.write(tf, meta, self.d["sizes"]["small"] + 3, "on_primary")
        self.footer(s, page, dark=True)

    def layout_section(self, s: Slide, page: int) -> None:
        from pptx.util import Inches

        self.new_slide(dark=True)
        tf = self.text_box(self.left, Inches(2.6), self.content_width(), Inches(1.4), anchor="middle")
        self.write(tf, s.title, self.d["sizes"]["title"], "on_primary", bold=True)
        if s.paragraphs:
            tf = self.text_box(self.left, Inches(4.1), self.content_width(), Inches(1.0))
            self.write(tf, s.paragraphs[0], self.d["sizes"]["body"], "on_primary")
        self.footer(s, page, dark=True)

    def bullets_into(self, tf, bullets: list[Bullet], paragraphs: list[str], size: float) -> None:
        from pptx.util import Pt

        first = True
        for text in paragraphs:
            p = self.write(tf, text, size, first=first)
            p.space_after = Pt(size * 0.5)
            first = False
        for b in bullets:
            marker = "• " if b.level == 0 else "– "
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.level = b.level
            p.space_after = Pt(size * 0.45)
            run = p.add_run()
            run.text = ("      " if b.level else "") + marker
            self.style_run(run, size if b.level == 0 else size - 2, "accent", bold=True)
            self.add_runs(p, b.text, size if b.level == 0 else size - 2)

    def layout_bullets(self, s: Slide, page: int) -> None:
        self.new_slide()
        y = self.heading(s)
        tf = self.text_box(self.left, y, self.content_width(), self.H - y - self.bottom - self.Inches(0.9))
        self.bullets_into(tf, s.bullets, s.paragraphs, self.d["sizes"]["body"])
        self.footer(s, page)

    def layout_two_column(self, s: Slide, page: int) -> None:
        from pptx.util import Inches

        self.new_slide()
        y = self.heading(s)
        gap = Inches(0.4)
        col_w = (self.content_width() - gap) / 2
        height = self.H - y - self.bottom - Inches(0.9)
        for i, (name, items) in enumerate(s.columns[:2]):
            x = self.left + i * (col_w + gap)
            self.rect(x, y, col_w, Inches(0.6), "light")
            tf = self.text_box(x + Inches(0.15), y, col_w - Inches(0.3), Inches(0.6), anchor="middle")
            self.write(tf, name, self.d["sizes"]["body"] + 2, "primary", bold=True)
            tf = self.text_box(x, y + Inches(0.75), col_w, height - Inches(0.75))
            self.bullets_into(tf, items, [], self.d["sizes"]["body"] - 1)
        self.footer(s, page)

    def layout_table(self, s: Slide, page: int) -> None:
        from pptx.util import Inches, Pt

        self.new_slide()
        y = self.heading(s)
        rows, cols = len(s.table), len(s.table[0])
        height = min(Inches(0.5) * rows, self.H - y - self.bottom - Inches(0.9))
        shape = self.slide.shapes.add_table(rows, cols, self.left, y, self.content_width(), height)
        table = shape.table
        # 列の幅は、各列の中身の最大の長さ（全角換算）に合わせて配分する。極端に狭くならないよう下限を設ける
        widths = [max(3.0, max(display_width(row[c]) if c < len(row) else 0 for row in s.table)) for c in range(cols)]
        total = sum(widths)
        for c in range(cols):
            table.columns[c].width = int(self.content_width() * widths[c] / total)
        size = self.d["sizes"]["body"] - 3
        for r, row in enumerate(s.table):
            for c in range(cols):
                cell = table.cell(r, c)
                cell.text = ""
                cell.fill.solid()
                cell.fill.fore_color.rgb = self.color("primary" if r == 0 else ("light" if r % 2 == 0 else "background"))
                cell.margin_left = cell.margin_right = Pt(6)
                text = row[c] if c < len(row) else ""
                self.add_runs(cell.text_frame.paragraphs[0], text, size,
                              "on_primary" if r == 0 else "text", bold=(r == 0))
        self.footer(s, page)

    def layout_stats(self, s: Slide, page: int) -> None:
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches

        self.new_slide()
        y = self.heading(s) + Inches(0.4)
        items = [STAT_RE.match(b.text) for b in s.bullets if STAT_RE.match(b.text)]
        n = max(len(items), 1)
        gap = Inches(0.3)
        card_w = (self.content_width() - gap * (n - 1)) / n
        card_h = Inches(2.6)
        for i, m in enumerate(items):
            x = self.left + i * (card_w + gap)
            self.rect(x, y, card_w, card_h, "light")
            self.rect(x, y, card_w, Inches(0.08), "accent")
            tf = self.text_box(x, y + Inches(0.3), card_w, Inches(1.1), anchor="middle")
            # 数値が長いときは、カードの幅に 1 行で収まるよう文字を小さくする（全角 1 文字 ≒ 文字サイズ 1pt 分の幅）
            usable_pt = (card_w - Inches(0.3)) / 12700
            size = min(self.d["sizes"]["stat"], usable_pt / max(display_width(m.group(1)), 1))
            self.write(tf, m.group(1), max(size, 14), "primary", bold=True, align=PP_ALIGN.CENTER)
            tf = self.text_box(x + Inches(0.2), y + Inches(1.45), card_w - Inches(0.4), Inches(1.0))
            self.write(tf, m.group(2), self.d["sizes"]["body"] - 2, "text", align=PP_ALIGN.CENTER)
        if s.paragraphs:
            tf = self.text_box(self.left, y + card_h + Inches(0.3), self.content_width(), Inches(0.9))
            self.write(tf, s.paragraphs[0], self.d["sizes"]["body"] - 2, "muted")
        self.footer(s, page)

    def layout_chart(self, s: Slide, page: int) -> None:
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
        from pptx.util import Inches, Pt

        self.new_slide()
        y = self.heading(s)
        data = CategoryChartData()
        data.categories = [row[0] for row in s.table[1:]]
        for c, name in enumerate(s.table[0][1:], start=1):
            data.add_series(name, [parse_number(row[c]) or 0 for row in s.table[1:]])
        kind = {"bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
                "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE}[s.chart or "column"]
        height = self.H - y - self.bottom - Inches(0.9)
        chart = self.slide.shapes.add_chart(kind, self.left, y, self.content_width(), height, data).chart
        chart.font.size = Pt(self.d["sizes"]["body"] - 4)
        chart.font.name = self.d["fonts"]["east_asian"]
        chart.has_legend = len(s.table[0]) > 2 or kind == XL_CHART_TYPE.PIE
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        palette = ["primary", "accent", "muted"]
        if kind != XL_CHART_TYPE.PIE:
            for i, series in enumerate(chart.series):
                fmt = series.format
                if kind == XL_CHART_TYPE.LINE_MARKERS:
                    fmt.line.color.rgb = self.color(palette[i % len(palette)])
                else:
                    fmt.fill.solid()
                    fmt.fill.fore_color.rgb = self.color(palette[i % len(palette)])
        self.footer(s, page)

    def layout_image(self, s: Slide, page: int) -> None:
        from pptx.util import Inches

        self.new_slide()
        y = self.heading(s) if s.title else self.top
        path = self.base_dir / s.image[1]
        max_w, max_h = self.content_width(), self.H - y - self.bottom - Inches(0.9)
        pic = self.slide.shapes.add_picture(str(path), self.left, y)
        scale = min(max_w / pic.width, max_h / pic.height, 1.0)
        pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
        pic.left = int(self.left + (max_w - pic.width) / 2)
        self.footer(s, page)

    def layout_message(self, s: Slide, page: int) -> None:
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches

        self.new_slide()
        self.rect(self.left, Inches(2.2), Inches(0.12), Inches(2.4), "accent")
        tf = self.text_box(self.left + Inches(0.4), Inches(2.0), self.content_width() - Inches(0.4), Inches(1.9),
                           anchor="middle")
        self.write(tf, s.title, self.d["sizes"]["message"], "primary", bold=True, align=PP_ALIGN.LEFT)
        if s.paragraphs:
            tf = self.text_box(self.left + Inches(0.4), Inches(4.0), self.content_width() - Inches(0.4), Inches(1.2))
            for i, text in enumerate(s.paragraphs):
                self.write(tf, text, self.d["sizes"]["body"], "muted", first=(i == 0))
        self.footer(s, page)

    def save(self, out: Path) -> None:
        out.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(out))


# --- 本体 -------------------------------------------------------------------------

def load_design(path: Path | None) -> dict:
    if path is None:
        return merge_design(DEFAULT_DESIGN, {})
    if not path.is_file():
        raise BuildError(f"デザインスペック {path} がない")
    import yaml

    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    if not isinstance(data, dict):
        raise BuildError(f"{path} は key: value の形にする")
    return merge_design(DEFAULT_DESIGN, data)


def main(argv: list[str]) -> int:
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")
    args, opts, i = [], {}, 0
    while i < len(argv):
        a = argv[i]
        if a in ("--design", "-o", "--output"):
            opts[a.lstrip("-")[0]] = argv[i + 1] if i + 1 < len(argv) else ""
            i += 2
            continue
        if a.startswith("--"):
            opts[a[2:]] = True
        else:
            args.append(a)
        i += 1
    if len(args) != 1:
        print(__doc__)
        return 2
    slides_path = Path(args[0]).resolve()
    try:
        deck = parse_deck(slides_path.read_text(encoding="utf-8"))
        base_dir = slides_path.parent
        if "d" in opts:
            design_path = Path(opts["d"]).resolve()
        elif deck.meta.get("design"):
            design_path = (base_dir / deck.meta["design"]).resolve()
        else:
            design_path = base_dir.parent / "design.yaml"
            if not design_path.is_file():
                design_path = None
        design = load_design(design_path)
        errors, warnings = check_deck(deck, design, base_dir)
        for w in warnings:
            print(f"[警告] {w}")
        for e in errors:
            print(f"[エラー] {e}")
        print(f"スライド {len(deck.slides)} 枚 / エラー {len(errors)} 件 / 警告 {len(warnings)} 件"
              f"（デザイン: {design_path or '既定値'}）")
        if errors or (opts.get("strict") and warnings):
            return 1
        if opts.get("check"):
            return 0
        out = Path(opts["o"]).resolve() if "o" in opts else base_dir / "proposal.pptx"
        renderer = Renderer(design, base_dir, design_path.parent if design_path else base_dir, deck.meta)
        renderer.render(deck)
        renderer.save(out)
        print(f"{out} を書き出した")
    except BuildError as error:
        print(f"[エラー] {error}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
